import docx
from docx.enum.text import WD_COLOR_INDEX

from app.ingest.docx_extractor import DocxExtractor


def _build_fixture(path):
    document = docx.Document()

    bold_paragraph = document.add_paragraph()
    bold_run = bold_paragraph.add_run("Bold text here.")
    bold_run.bold = True

    highlight_paragraph = document.add_paragraph()
    highlight_run = highlight_paragraph.add_run("Highlighted text here.")
    highlight_run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    comment_paragraph = document.add_paragraph()
    comment_run = comment_paragraph.add_run("Text with a comment.")
    document.add_comment(runs=comment_run, text="Author note here", author="Author")

    document.save(path)


def test_docx_extractor_captures_formatting_and_comments(tmp_path):
    fixture_path = tmp_path / "sample.docx"
    _build_fixture(fixture_path)

    result = DocxExtractor().extract(fixture_path)

    assert result.source_type == "docx"

    segments_by_text = {segment.text: segment for segment in result.segments}

    bold_styles = {s.style_kind for s in segments_by_text["Bold text here."].styles}
    assert "bold" in bold_styles

    highlight_segment = segments_by_text["Highlighted text here."]
    highlight_styles = {s.style_kind: s.style_value for s in highlight_segment.styles}
    assert highlight_styles.get("highlight") == "YELLOW"

    comment_segment = segments_by_text["Text with a comment."]
    comment_styles = {s.style_kind: s.style_value for s in comment_segment.styles}
    assert comment_styles.get("comment") == "Author note here"


def test_docx_extractor_sequence_index_increments(tmp_path):
    fixture_path = tmp_path / "sample.docx"
    _build_fixture(fixture_path)

    result = DocxExtractor().extract(fixture_path)

    assert [s.sequence_index for s in result.segments] == list(range(len(result.segments)))


def test_docx_extractor_skips_whitespace_only_runs(tmp_path):
    document = docx.Document()
    document.add_paragraph().add_run("Real text.")
    document.add_paragraph().add_run("          ")
    document.save(tmp_path / "sample.docx")

    result = DocxExtractor().extract(tmp_path / "sample.docx")

    assert [s.text for s in result.segments] == ["Real text."]


def test_docx_extractor_tags_heading_one_paragraphs(tmp_path):
    document = docx.Document()
    document.add_paragraph("Chapter One", style="Heading 1")
    document.add_paragraph().add_run("Body text.")
    document.save(tmp_path / "sample.docx")

    result = DocxExtractor().extract(tmp_path / "sample.docx")

    segments_by_text = {segment.text: segment for segment in result.segments}
    heading_styles = {s.style_kind for s in segments_by_text["Chapter One"].styles}
    body_styles = {s.style_kind for s in segments_by_text["Body text."].styles}
    assert "heading" in heading_styles
    assert "heading" not in body_styles
