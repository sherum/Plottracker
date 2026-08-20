import docx
import pymupdf
from docx.enum.text import WD_COLOR_INDEX

from app.db import repository
from app.ingest.service import ingest_folder


def _build_fixtures(folder):
    (folder / "notes.txt").write_text("First paragraph.\n\nSecond paragraph.")

    document = docx.Document()
    bold_run = document.add_paragraph().add_run("Bold text here.")
    bold_run.bold = True
    highlight_run = document.add_paragraph().add_run("Highlighted text here.")
    highlight_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    comment_run = document.add_paragraph().add_run("Text with a comment.")
    document.add_comment(runs=comment_run, text="Author note", author="Author")
    document.save(folder / "chapter1.docx")

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Bold pdf text.", fontname="hebo", fontsize=12)
    page.insert_text((72, 100), "Highlighted pdf text.", fontname="helv", fontsize=12)
    page.add_highlight_annot(pymupdf.Rect(65, 85, 220, 105))
    pdf.save(folder / "chapter2.pdf")


def test_ingest_folder_round_trip(db_conn, tmp_path):
    fixtures_dir = tmp_path / "fixtures"
    fixtures_dir.mkdir()
    _build_fixtures(fixtures_dir)

    summary = ingest_folder(db_conn, fixtures_dir, role="draft_script")

    assert set(summary["ingested"]) == {"notes.txt", "chapter1.docx", "chapter2.pdf"}
    assert summary["skipped"] == []
    assert summary["failed"] == []

    documents = repository.list_documents(db_conn)
    assert len(documents) == 3
    assert {d["source_type"] for d in documents} == {"txt", "docx", "pdf"}

    all_styles = []
    for document in documents:
        for segment in repository.get_segments(db_conn, document["id"]):
            all_styles.extend(s["style_kind"] for s in segment["styles"])

    assert "bold" in all_styles
    assert "highlight" in all_styles
    assert "comment" in all_styles


def test_ingest_folder_missing_directory_raises(db_conn, tmp_path):
    try:
        ingest_folder(db_conn, tmp_path / "does_not_exist", role="draft_script")
        assert False, "expected ValueError"
    except ValueError:
        pass


def test_ingest_folder_skips_unsupported_files(db_conn, tmp_path):
    fixtures_dir = tmp_path / "fixtures"
    fixtures_dir.mkdir()
    (fixtures_dir / "cover.png").write_bytes(b"not a real png")

    summary = ingest_folder(db_conn, fixtures_dir, role="story_note")

    assert summary["ingested"] == []
    assert summary["skipped"] == ["cover.png"]
    assert summary["failed"] == []
