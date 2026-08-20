import docx


def _build_docx_with_dream_sequence(folder):
    document = docx.Document()
    document.add_paragraph("Chapter One", style="Heading 1")
    document.add_paragraph()
    # A whitespace-only paragraph between the heading and the first real
    # content is common in real manuscripts (leftover indentation/spacing)
    # and must not break chapter-start detection.
    document.add_paragraph().add_run("   ")

    document.add_paragraph().add_run("I was flying.").italic = True
    document.add_paragraph().add_run("The ground fell away.").italic = True
    document.add_paragraph("I woke up in my bed.")
    document.add_paragraph().add_run("Where am I really?").italic = True
    document.add_paragraph("He kept walking.")
    document.save(folder / "chapter.docx")


def test_full_pipeline_classifies_a_real_docx(client, tmp_path):
    fixtures_dir = tmp_path / "fixtures"
    fixtures_dir.mkdir()
    _build_docx_with_dream_sequence(fixtures_dir)

    client.post("/ingest", json={"folder_path": str(fixtures_dir), "role": "draft_script"})
    document_id = client.get("/documents").json()[0]["id"]

    client.post(
        "/encoding-rules",
        json={"style_kind": "italic", "block_length": "multi", "position": "chapter_start", "label": "dream_sequence"},
    )
    client.post(
        "/encoding-rules",
        json={"style_kind": "italic", "block_length": "single", "position": "anywhere", "label": "internal_dialogue"},
    )
    assert len(client.get("/encoding-rules").json()) == 2

    result = client.post(f"/documents/{document_id}/classify-encoding")
    assert result.status_code == 200
    assert result.json() == {"tagged": 3}

    segments = client.get(f"/documents/{document_id}/segments").json()
    tags = {
        s["text"]: [st["style_value"] for st in s["styles"] if st["style_kind"] == "semantic"] for s in segments
    }
    assert tags["I was flying."] == ["dream_sequence"]
    assert tags["The ground fell away."] == ["dream_sequence"]
    assert tags["Where am I really?"] == ["internal_dialogue"]
    assert tags["I woke up in my bed."] == []
