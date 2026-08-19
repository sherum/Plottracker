from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.text.run import Run

from app.ingest.models import ExtractedDocument, ExtractedSegment, ExtractedStyle

COMMENT_RANGE_START = qn("w:commentRangeStart")
COMMENT_RANGE_END = qn("w:commentRangeEnd")
COMMENT_ID_ATTR = qn("w:id")
RUN_TAG = qn("w:r")


class DocxExtractor:
    def extract(self, path: Path) -> ExtractedDocument:
        document = docx.Document(path)
        comments_by_id = {c.comment_id: c for c in document.comments}

        segments: list[ExtractedSegment] = []
        sequence_index = 0
        active_comment_ids: set[int] = set()

        for paragraph_index, paragraph in enumerate(document.paragraphs):
            for element in paragraph._p:
                if element.tag == COMMENT_RANGE_START:
                    active_comment_ids.add(int(element.get(COMMENT_ID_ATTR)))
                elif element.tag == COMMENT_RANGE_END:
                    active_comment_ids.discard(int(element.get(COMMENT_ID_ATTR)))
                elif element.tag == RUN_TAG:
                    run = Run(element, paragraph)
                    if not run.text:
                        continue

                    styles = self._run_styles(run, active_comment_ids, comments_by_id)

                    segments.append(
                        ExtractedSegment(
                            sequence_index=sequence_index,
                            paragraph_index=paragraph_index,
                            text=run.text,
                            styles=styles,
                        )
                    )
                    sequence_index += 1

        return ExtractedDocument(source_type="docx", segments=segments)

    @staticmethod
    def _run_styles(run: Run, active_comment_ids: set[int], comments_by_id: dict) -> list[ExtractedStyle]:
        styles: list[ExtractedStyle] = []

        if run.bold:
            styles.append(ExtractedStyle(style_kind="bold"))
        if run.italic:
            styles.append(ExtractedStyle(style_kind="italic"))
        if run.underline:
            styles.append(ExtractedStyle(style_kind="underline"))
        if run.font.highlight_color is not None:
            styles.append(ExtractedStyle(style_kind="highlight", style_value=run.font.highlight_color.name))
        for comment_id in active_comment_ids:
            comment = comments_by_id.get(comment_id)
            if comment is not None:
                styles.append(ExtractedStyle(style_kind="comment", style_value=comment.text))

        return styles
